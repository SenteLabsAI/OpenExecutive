"""Delivering artifacts to people.

Two paths, both keyed on an artifact id and both resolved through
`orchestrator.artifact_records` (so only artifact rows can ever be shared):

- `message_person(..., artifact_id=...)` appends the artifact's title and a
  `UI_BASE_URL/artifacts/<id>` link to the DM.
- a Gmail send/draft through the MCP gateway may attach
  `{"artifact_id": ..., "as"?: ...}`; the gateway renders the file exactly as
  `/artifacts/<id>/download` serves it — only after the recipient gate passes.
"""
from __future__ import annotations

import asyncio
import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from openpyxl import load_workbook

import openexecutive.orchestrator.mcp_gateway as gw_module
from openexecutive.alerts import store as alerts_store
from openexecutive.orchestrator.mcp_gateway import MCPGateway
from openexecutive.orchestrator.schedule_tools import handle_message_person
from openexecutive.people import store as people_store
from openexecutive.workflows import persistence as wf_persistence

EXEC_ADDR = "exec@example.com"
UI = "https://oe.example.com/"


@pytest.fixture(autouse=True)
def db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db_path = tmp_path / "episodic.db"
    for module in (alerts_store, people_store, wf_persistence):
        monkeypatch.setattr(module, "DB_PATH", db_path)
    alerts_store.initialize_db(db_path)
    people_store.initialize_db()
    wf_persistence.initialize_runs_db(db_path)
    monkeypatch.setattr("openexecutive.audit.log_event", lambda *a, **k: None)
    return db_path


def _artifact(fmt: str = "markdown", body: str = "# Q3 plan\n\nHire two engineers.",
              title: str = "Q3 plan", **extra: Any) -> str:
    aid = alerts_store.insert_alert(
        source="artifact", external_id=f"{fmt}-{title}", severity="medium",
        headline=title, body=body, topic_tags=["artifact"],
        artifact_format=fmt, **extra,
    )
    assert aid is not None
    return f"alert:{aid}"


# --------------------------------------------------------------------------- #
# message_person(artifact_id=...)
# --------------------------------------------------------------------------- #


def _dm_settings() -> SimpleNamespace:
    return SimpleNamespace(
        slack_bot_token=None, discord_bot_token="tok", telegram_bot_token=None,
        calendar_booking_enabled=False, mcp_enabled=False, ui_base_url=UI,
    )


def _dm(payload: dict) -> tuple[dict, AsyncMock]:
    send = AsyncMock(return_value=None)
    with (
        patch("openexecutive.config.get_settings", return_value=_dm_settings()),
        patch("openexecutive.integrations.discord_bot.send_dm", send),
        patch("openexecutive.orchestrator.schedule_tools._record_send_to_activity"),
    ):
        result = json.loads(asyncio.run(handle_message_person(payload)))
    return result, send


def _person() -> int:
    return people_store.upsert_person(
        full_name="Dana Ops", discord_user_id="100000000000000001",
        preferred_channel="discord",
    )


def test_message_person_appends_artifact_title_and_link() -> None:
    cid = _artifact(title="Q3   hiring\nplan")
    result, send = _dm({"person_id": _person(), "text": "Here's the plan.",
                        "artifact_id": cid})
    assert result["status"] == "sent"
    sent_text = send.await_args.args[1]
    assert sent_text == (
        f"Here's the plan.\n\n📄 Q3 hiring plan — "
        f"https://oe.example.com/artifacts/{cid}"
    )


def test_message_person_links_workflow_run_artifact() -> None:
    wf_persistence.create_run("run-7", "board_prep", "Board deck", {})
    wf_persistence.complete_run("run-7", "# Deck")
    _, send = _dm({"person_id": _person(), "text": "Deck", "artifact_id": "run:run-7"})
    assert send.await_args.args[1].endswith("https://oe.example.com/artifacts/run:run-7")


@pytest.mark.parametrize("bad", ["alert:999", "garbage", "alert:x"])
def test_message_person_rejects_unknown_artifact_and_sends_nothing(bad: str) -> None:
    result, send = _dm({"person_id": _person(), "text": "hi", "artifact_id": bad})
    assert "artifact_id" in result["error"]
    send.assert_not_awaited()


def test_message_person_refuses_non_artifact_alert() -> None:
    other = alerts_store.insert_alert(source="email", external_id="e", severity="low",
                                      headline="Private", body="secret")
    result, send = _dm({"person_id": _person(), "text": "hi",
                        "artifact_id": f"alert:{other}"})
    assert "error" in result
    send.assert_not_awaited()


def test_message_person_without_artifact_is_unchanged() -> None:
    result, send = _dm({"person_id": _person(), "text": "plain"})
    assert result["status"] == "sent"
    assert send.await_args.args[1] == "plain"


# --------------------------------------------------------------------------- #
# Gmail attachments through the MCP gateway
# --------------------------------------------------------------------------- #


def _gateway() -> tuple[MCPGateway, AsyncMock]:
    gateway = MCPGateway()
    session = MagicMock()
    fake_result = MagicMock()
    fake_result.content = [MagicMock(text='{"ok": true}')]
    session.call_tool = AsyncMock(return_value=fake_result)
    gateway._session = session
    return gateway, session.call_tool


def _send(arguments: dict[str, Any],
          tool: str = "google_workspace__send_gmail_message") -> tuple[str, AsyncMock]:
    people_store.upsert_person(full_name="Alice", email="alice@example.com")
    gateway, session_call = _gateway()
    settings = SimpleNamespace(exec_email_address=EXEC_ADDR, email_poll_interval_seconds=60)
    with patch.object(gw_module, "get_settings", return_value=settings):
        result = asyncio.run(gateway.call_tool({"name": tool, "arguments": arguments}))
    return result, session_call


def _forwarded_attachments(session_call: AsyncMock) -> list[dict]:
    assert session_call.await_count == 1
    forwarded = session_call.await_args.args[1]
    return forwarded["arguments"]["attachments"]


@pytest.mark.parametrize("tool", [
    "google_workspace__send_gmail_message", "google_workspace__draft_gmail_message",
])
def test_gmail_attaches_artifact_as_its_own_format(tool: str) -> None:
    cid = _artifact()
    result, session_call = _send({"to": "alice@example.com", "subject": "Plan",
                                  "body": "Attached.", "attachments": [{"artifact_id": cid}]},
                                 tool=tool)
    assert result == '{"ok": true}'
    (att,) = _forwarded_attachments(session_call)
    assert att["filename"] == "q3-plan.md"
    assert att["mime_type"].startswith("text/markdown")
    assert base64.b64decode(att["content"]).decode() == "# Q3 plan\n\nHire two engineers."


def test_gmail_attaches_markdown_artifact_as_word() -> None:
    cid = _artifact()
    _, session_call = _send({"to": "alice@example.com", "subject": "Plan", "body": "x",
                             "attachments": [{"artifact_id": cid, "as": "docx"}]})
    (att,) = _forwarded_attachments(session_call)
    assert att["filename"] == "q3-plan.docx"
    doc = Document(io.BytesIO(base64.b64decode(att["content"])))
    assert any(p.text == "Hire two engineers." for p in doc.paragraphs)


def test_gmail_attaches_spreadsheet_and_keeps_other_attachments() -> None:
    cid = _artifact("xlsx", json.dumps({"summary": "", "sheets": [
        {"name": "Tiers", "columns": ["tier"], "rows": [["pro"]]}]}), title="Pricing")
    passthrough = {"path": "/tmp/already-downloaded.pdf"}
    _, session_call = _send({"to": "alice@example.com", "subject": "s", "body": "b",
                             "attachments": [passthrough, {"artifact_id": cid}]})
    first, second = _forwarded_attachments(session_call)
    assert first == passthrough
    wb = load_workbook(io.BytesIO(base64.b64decode(second["content"])))
    assert wb["Tiers"]["A2"].value == "pro"


def test_blocked_recipient_never_renders_the_artifact(monkeypatch: pytest.MonkeyPatch) -> None:
    rendered: list[str] = []
    import openexecutive.orchestrator.artifact_records as records

    real = records.render_artifact_file

    def _spy(rec, as_=None):  # noqa: ANN001, ANN202
        rendered.append(rec.id)
        return real(rec, as_)

    monkeypatch.setattr(records, "render_artifact_file", _spy)
    cid = _artifact()
    result, session_call = _send({"to": "attacker@evil.com", "subject": "s", "body": "b",
                                  "attachments": [{"artifact_id": cid}]})
    assert "error" in json.loads(result)
    assert session_call.await_count == 0
    assert rendered == []


@pytest.mark.parametrize("entry, needle", [
    ({"artifact_id": "alert:4242"}, "not found"),
    ({"artifact_id": "nonsense"}, "Malformed"),
    ({"artifact_id": "ALERT_PLACEHOLDER", "as": "xlsx"}, "no xlsx download"),
    ({"artifact_id": "ALERT_PLACEHOLDER", "path": "/etc/passwd"}, "only 'artifact_id'"),
])
def test_gmail_artifact_attachment_errors_send_nothing(entry: dict, needle: str) -> None:
    cid = _artifact()
    entry = {k: (cid if v == "ALERT_PLACEHOLDER" else v) for k, v in entry.items()}
    result, session_call = _send({"to": "alice@example.com", "subject": "s", "body": "b",
                                  "attachments": [entry]})
    assert needle in json.loads(result)["error"]
    assert session_call.await_count == 0


def test_gmail_refuses_non_artifact_alert_and_links() -> None:
    other = alerts_store.insert_alert(source="email", external_id="e", severity="low",
                                      headline="Private", body="secret")
    link = _artifact("link", "Summary", title="Sheet", artifact_url="https://x.example/s")
    for cid in (f"alert:{other}", link):
        result, session_call = _send({"to": "alice@example.com", "subject": "s",
                                      "body": "b", "attachments": [{"artifact_id": cid}]})
        assert "error" in json.loads(result)
        assert session_call.await_count == 0


def test_gmail_rejects_oversized_artifact(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(gw_module, "_MAX_ARTIFACT_ATTACHMENT_BYTES", 10)
    cid = _artifact()
    result, session_call = _send({"to": "alice@example.com", "subject": "s", "body": "b",
                                  "attachments": [{"artifact_id": cid}]})
    assert "email limit" in json.loads(result)["error"]
    assert session_call.await_count == 0


def test_gmail_without_artifact_attachments_is_untouched() -> None:
    args = {"to": "alice@example.com", "subject": "s", "body": "b",
            "attachments": [{"path": "/tmp/x.pdf"}]}
    _, session_call = _send(dict(args))
    assert session_call.await_args.args[1]["arguments"] == args
