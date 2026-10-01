"""knowledge.isolated — document parsing in a short-lived child process.

These run the real child (``python -m openexecutive.knowledge.isolated``):
the reads must come back exactly as they do in-process, and a child that
fails, dies or hangs must surface as an error the readers already handle,
never as a hung or crashed API process.
"""
from __future__ import annotations

import math
import os
import time
from pathlib import Path

import pytest

from openexecutive.knowledge import isolated, pdf_reader
from openexecutive.knowledge.isolated import IsolatedError, WorkerStopped, run_isolated
from openexecutive.knowledge.loader import _parse_file, extract_text_from_file

from .test_pdf_reader import _blank_pdf, _image_pdf, _text_pdf


@pytest.fixture(autouse=True)
def _fresh_cache():
    pdf_reader.clear_cache()
    pdf_reader.reset_inbound_budget()
    yield
    pdf_reader.clear_cache()
    pdf_reader.reset_inbound_budget()


# ── The mechanism ────────────────────────────────────────────────────────────


def test_a_result_comes_back_from_the_child() -> None:
    assert run_isolated(math.sqrt, 16.0, timeout=60) == 4.0


def test_the_child_does_not_inherit_the_deployments_keys(monkeypatch) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-should-not-leak")

    assert run_isolated(os.getenv, "ANTHROPIC_API_KEY", timeout=60) is None


def test_a_named_exception_is_raised_again_as_itself() -> None:
    with pytest.raises(ValueError, match="math domain error"):
        run_isolated(math.sqrt, -1.0, timeout=60, reraise=(ValueError,))


def test_any_other_exception_is_an_isolated_error() -> None:
    with pytest.raises(IsolatedError, match="ValueError: math domain error") as info:
        run_isolated(math.sqrt, -1.0, timeout=60)
    assert not isinstance(info.value, WorkerStopped)


def test_a_child_that_dies_without_answering_is_reported() -> None:
    with pytest.raises(WorkerStopped, match="exit code 3"):
        run_isolated(os._exit, 3, timeout=60)


def test_a_child_past_its_timeout_is_killed() -> None:
    started = time.monotonic()
    with pytest.raises(WorkerStopped, match="timed out"):
        run_isolated(time.sleep, 30, timeout=1)
    assert time.monotonic() - started < 20


def test_without_a_child_process_it_parses_in_process(monkeypatch) -> None:
    monkeypatch.setattr(isolated.sys, "executable", "/nonexistent/python")

    assert run_isolated(math.sqrt, 9.0, timeout=60) == 3.0


# ── PDFs: same reads as in-process ───────────────────────────────────────────


async def test_a_text_layer_pdf_reads_the_same_as_in_process(tmp_path: Path, monkeypatch) -> None:
    data = _text_pdf("Revenue grew 30 percent year over year in every region we serve")
    path = tmp_path / "report.pdf"
    path.write_bytes(data)

    from_path = await pdf_reader.read_pdf_text(path, filename="report.pdf")
    pdf_reader.clear_cache()
    from_bytes = await pdf_reader.read_pdf_text(data, filename="report.pdf")
    pdf_reader.clear_cache()
    monkeypatch.setattr(isolated, "enabled", False)
    in_process = await pdf_reader.read_pdf_text(data, filename="report.pdf")

    assert from_path == from_bytes == in_process
    assert from_path.method == "text_layer"
    assert "Revenue grew 30 percent" in from_path.text


async def test_a_damaged_pdf_is_reported_through_the_child(tmp_path: Path) -> None:
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"%PDF-1.4 not really a pdf")

    result = await pdf_reader.read_pdf_text(path, filename="bad.pdf")

    assert result.method == "none"
    assert "could not be opened" in result.note


async def test_the_page_ceiling_holds_in_the_child(monkeypatch) -> None:
    """The ceiling is passed to the child, which never sees a patched module."""
    monkeypatch.setattr(pdf_reader, "_MAX_PDF_PAGES", 5)

    result = await pdf_reader.read_pdf_text(_blank_pdf(8), filename="huge.pdf")

    assert result == pdf_reader.PdfReadResult(
        "", "none", 8, "the PDF has 8 pages — more than the 5 this reads"
    )


async def test_a_reader_that_runs_out_of_time_never_raises(monkeypatch) -> None:
    monkeypatch.setattr(pdf_reader, "_TEXT_LAYER_TIMEOUT_S", 0.001)

    result = await pdf_reader.read_pdf_text(_text_pdf("Quarterly plan"), filename="slow.pdf")

    assert result.method == "none"
    assert result.note == "the PDF could not be read: reading it ran out of memory or time"


async def test_a_missing_file_is_reported_not_raised(tmp_path: Path) -> None:
    result = await pdf_reader.read_pdf_text(tmp_path / "gone.pdf", filename="gone.pdf")

    assert result.method == "none"
    assert result.note == "the file could not be read"


async def test_a_scan_is_ocrd_in_the_child(tmp_path: Path, monkeypatch) -> None:
    pytest.importorskip("rapidocr_onnxruntime")
    pytest.importorskip("pypdfium2")
    monkeypatch.setenv("PDF_PROVIDER_READING", "false")
    monkeypatch.setenv("PDF_OCR_ENABLED", "true")
    path = tmp_path / "scan.pdf"
    path.write_bytes(_image_pdf(["Board meeting notes", "Approve the hiring plan"]))

    result = await pdf_reader.read_pdf_text(path, filename="scan.pdf")

    assert result.method == "ocr"
    assert "Board meeting notes" in result.text
    assert "Approve the hiring plan" in result.text
    # The OCR model was loaded in the child, not here.
    assert pdf_reader._ocr_engine is None


# ── Word / Excel ─────────────────────────────────────────────────────────────


def test_a_docx_reads_the_same_as_in_process(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hiring plan")
    doc.add_paragraph("")
    doc.add_paragraph("Two engineers in Q3, one designer in Q4.")
    path = tmp_path / "plan.docx"
    doc.save(str(path))

    text = extract_text_from_file(path)

    assert text == _parse_file(path)
    assert text == "Hiring plan\n\nTwo engineers in Q3, one designer in Q4."


def test_a_corrupt_docx_raises_naming_the_parser_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")

    with pytest.raises(IsolatedError, match="PackageNotFoundError"):
        extract_text_from_file(path)


def test_what_a_parser_prints_does_not_spoil_the_answer() -> None:
    """Output a library writes to stdout goes to stderr, not into the JSON."""
    assert run_isolated(print, "noise from a parser", timeout=60) is None
