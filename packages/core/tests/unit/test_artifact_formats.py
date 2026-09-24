"""Unit tests for the artifact format registry (orchestrator/artifact_formats.py)."""
from __future__ import annotations

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook

from openexecutive.orchestrator import artifact_formats as af
from openexecutive.orchestrator.artifact_tools import DRAFT_ARTIFACT_TOOL


def test_format_names_are_sorted_and_match_registry() -> None:
    # The tuple feeds the prompt-cached tool schema; it must be stable.
    assert list(af.ARTIFACT_FORMAT_NAMES) == sorted(af.ARTIFACT_FORMAT_NAMES)
    assert set(af.ARTIFACT_FORMAT_NAMES) == set(af.ARTIFACT_FORMATS)
    enum = DRAFT_ARTIFACT_TOOL["input_schema"]["properties"]["format"]["enum"]
    assert enum == list(af.ARTIFACT_FORMAT_NAMES)


def test_tool_schema_serializes_identically() -> None:
    # Byte-stable across calls: no sets, no iteration-order dependence.
    first = json.dumps(DRAFT_ARTIFACT_TOOL, sort_keys=False)
    second = json.dumps(DRAFT_ARTIFACT_TOOL, sort_keys=False)
    assert first == second
    assert DRAFT_ARTIFACT_TOOL["input_schema"]["required"] == ["title", "why_interesting"]


def test_markdown_build_appends_sources() -> None:
    built = af.build_artifact("markdown", {
        "document": "# Memo\n\nBody.", "source_urls": ["https://a.example", " "],
    })
    assert built.stored.startswith("# Memo")
    assert "### Sources\n- https://a.example" in built.stored
    assert built.url is None


def test_markdown_requires_document() -> None:
    with pytest.raises(af.ArtifactInputError):
        af.build_artifact("markdown", {"document": "   "})


def test_document_cap_enforced() -> None:
    with pytest.raises(af.ArtifactInputError, match="limit"):
        af.build_artifact("markdown", {"document": "x" * (af.MAX_DOCUMENT_CHARS + 1)})


def test_unknown_format_rejected() -> None:
    with pytest.raises(af.ArtifactInputError, match="unknown format"):
        af.build_artifact("pdf", {"document": "x"})


def test_html_strips_active_content() -> None:
    html = (
        "<html><head><base href='https://evil.example/'>"
        "<meta http-equiv=\"refresh\" content=\"0;url=https://evil.example\">"
        "<style>h1{color:red}</style></head>"
        "<body onload=\"steal()\"><h1>Report</h1><script>alert(1)</script>"
        "<p onclick='x()'>Revenue up 12%.</p><img src=x onerror=alert(1)></body></html>"
    )
    stored = af.build_artifact("html", {"document": html}).stored
    lowered = stored.lower()
    for bad in ("<script", "<base", "refresh", "onload", "onclick", "onerror"):
        assert bad not in lowered
    assert "<h1>Report</h1>" in stored
    assert "color:red" in stored  # inline styling survives


def test_html_text_is_readable_and_tag_free() -> None:
    text = af.html_to_text(
        "<style>.a{}</style><h1>Title</h1><p>First &amp; second.</p><ul><li>One</li></ul>"
    )
    assert "<" not in text and ".a{}" not in text
    assert "Title" in text and "First & second." in text and "One" in text


def test_html_without_visible_text_rejected() -> None:
    with pytest.raises(af.ArtifactInputError, match="no visible text"):
        af.build_artifact("html", {"document": "<script>alert(1)</script>"})


def test_docx_renders_headings_lists_tables() -> None:
    md = (
        "# Board memo\n\nIntro with **bold** and *italic*.\n\n"
        "## Numbers\n\n- first point\n- second point\n\n1. step one\n\n"
        "| Metric | Q1 |\n|---|---|\n| ARR | 1.2M |\n"
    )
    built = af.build_artifact("docx", {"document": md})
    assert built.stored == md.strip()  # Markdown is the stored source

    doc = Document(io.BytesIO(af.markdown_to_docx(built.stored)))
    texts = [p.text for p in doc.paragraphs]
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Board memo"] == "Heading 1"
    assert styles["Numbers"] == "Heading 2"
    assert styles["first point"] == "List Bullet"
    assert styles["step one"] == "List Number"
    assert "Intro with bold and italic." in texts
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "bold" in bold_runs
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Metric", "Q1"]
    assert [c.text for c in table.rows[1].cells] == ["ARR", "1.2M"]


def test_docx_keeps_unrecognised_text() -> None:
    doc = Document(io.BytesIO(af.markdown_to_docx("```\ncode line\n```\n\n> quoted\n\nplain")))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "code line" in joined and "quoted" in joined and "plain" in joined


def _sheets() -> list[dict]:
    return [
        {"name": "Pricing", "columns": ["Vendor", "Price", "Notes"],
         "rows": [["Acme", 12.5, "=HYPERLINK(\"https://evil\")"], ["Beta", 9, None]]},
        {"name": "Pricing", "columns": ["X"], "rows": [[1]]},
    ]


def test_xlsx_build_and_render() -> None:
    built = af.build_artifact("xlsx", {"document": "Two vendors compared.", "sheets": _sheets()})
    data = json.loads(built.stored)
    assert data["summary"] == "Two vendors compared."
    # Duplicate sheet names are made unique.
    assert [s["name"] for s in data["sheets"]] == ["Pricing", "Pricing (2)"]

    wb = load_workbook(io.BytesIO(af.sheets_to_xlsx(built.stored)))
    assert wb.sheetnames == ["Pricing", "Pricing (2)"]
    ws = wb["Pricing"]
    assert [c.value for c in ws[1]] == ["Vendor", "Price", "Notes"]
    assert ws["A2"].value == "Acme" and ws["B2"].value == 12.5
    # Text that looks like a formula is stored as text, never as a formula.
    assert ws["C2"].data_type == "s"
    assert ws["C2"].value == "=HYPERLINK(\"https://evil\")"


def test_xlsx_display_is_markdown_table_with_row_cap() -> None:
    rows = [[i, f"row {i}"] for i in range(af.PREVIEW_ROWS + 5)]
    built = af.build_artifact("xlsx", {"sheets": [{"name": "Big", "columns": ["n", "label"],
                                                   "rows": rows}]})
    shown = af.get_format("xlsx").display(built.stored)
    assert "### Big" in shown
    assert "| n | label |" in shown
    assert f"Showing {af.PREVIEW_ROWS} of {af.PREVIEW_ROWS + 5} rows" in shown
    assert f"row {af.PREVIEW_ROWS + 1}" not in shown


@pytest.mark.parametrize("sheets", [
    None, [], [{"columns": [], "rows": []}], "nope",
    [{"columns": ["a"], "rows": [[1]]}] * (af.MAX_SHEETS + 1),
    [{"columns": ["a"], "rows": [[1]] * (af.MAX_ROWS_PER_SHEET + 1)}],
])
def test_xlsx_rejects_bad_sheets(sheets: object) -> None:
    with pytest.raises(af.ArtifactInputError):
        af.build_artifact("xlsx", {"sheets": sheets})


def test_link_build() -> None:
    built = af.build_artifact("link", {
        "url": "https://docs.example.com/sheet/1", "link_label": "  Google   Sheet ",
        "document": "Pricing model with three scenarios.",
    })
    assert built.url == "https://docs.example.com/sheet/1"
    assert built.link_label == "Google Sheet"
    assert built.stored == "Pricing model with three scenarios."


def test_link_label_defaults() -> None:
    built = af.build_artifact("link", {"url": "https://notion.example/p"})
    assert built.link_label == "External document"


@pytest.mark.parametrize("url", [
    "", "http://example.com", "javascript:alert(1)", "https://",
    "https://user:pw@example.com/x", "https://exa mple.com", "ftp://example.com",
])
def test_link_rejects_unsafe_urls(url: str) -> None:
    with pytest.raises(af.ArtifactInputError):
        af.build_artifact("link", {"url": url})


def test_get_format_falls_back_to_markdown() -> None:
    assert af.get_format(None).name == "markdown"
    assert af.get_format("legacy-thing").name == "markdown"
    assert af.get_format(" HTML ").name == "html"


def test_docx_keeps_intraword_underscores() -> None:
    doc = Document(io.BytesIO(af.markdown_to_docx("Use snake_case_name and john_doe_x@example.com, _emphasis_ too.")))
    text = doc.paragraphs[0].text
    assert "snake_case_name" in text and "john_doe_x@example.com" in text
    assert "emphasis too." in text
    assert any(r.italic and r.text == "emphasis" for r in doc.paragraphs[0].runs)


def test_docx_table_keeps_extra_cells_and_ignores_bare_rule() -> None:
    doc = Document(io.BytesIO(af.markdown_to_docx("| a | b |\n|---|---|\n| 1 | 2 | 3 |")))
    table = doc.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["1", "2", "3"]
    # A rule under a line with a pipe is not a table (delimiter width differs).
    doc2 = Document(io.BytesIO(af.markdown_to_docx("Revenue | margin notes\n---\nnext para")))
    assert doc2.tables == []
    assert "Revenue | margin notes" in "\n".join(p.text for p in doc2.paragraphs)


def test_control_characters_never_break_rendering() -> None:
    af.markdown_to_docx("a\x01b\x00c\x0b")  # would raise inside python-docx
    built = af.build_artifact("xlsx", {"sheets": [
        {"name": "S\x01", "columns": ["c\x02"], "rows": [["x\x0by", 1]]},
    ]})
    wb = load_workbook(io.BytesIO(af.sheets_to_xlsx(built.stored)))
    ws = wb[wb.sheetnames[0]]
    assert ws["A1"].value == "c" and ws["A2"].value == "xy"
    # Legacy stored JSON that still holds control characters renders too.
    legacy = json.dumps({"summary": "", "sheets": [{"name": "L", "columns": ["a"], "rows": [["\x01z"]]}]})
    assert load_workbook(io.BytesIO(af.sheets_to_xlsx(legacy)))["L"]["A2"].value == "z"


def test_xlsx_accepts_rows_as_objects() -> None:
    built = af.build_artifact("xlsx", {"sheets": [
        {"name": "Deals", "rows": [{"Name": "Acme", "ARR": 10}, {"ARR": 5, "Name": "Beta"}]},
    ]})
    sheet = json.loads(built.stored)["sheets"][0]
    assert sheet["columns"] == ["Name", "ARR"]
    assert sheet["rows"] == [["Acme", 10], ["Beta", 5]]


def test_link_rejects_unparseable_url() -> None:
    with pytest.raises(af.ArtifactInputError, match="not valid"):
        af.build_artifact("link", {"url": "https://[abc/x"})


def test_heading_only_markdown_keeps_preview() -> None:
    assert af.preview_text("markdown", "# Title only", 200) == "# Title only"


def test_html_sanitizer_keeps_body_text_and_catches_slash_handlers() -> None:
    out = af.sanitize_html(
        '<p>Set online = true today</p><img/onerror=alert(1) src=x>'
        '<a href="javascript:alert(1)">x</a>'
    )
    assert "Set online = true today" in out
    assert "onerror" not in out.lower()
    assert "javascript:" not in out.lower()


@pytest.mark.parametrize("payload", [
    '<meta http-equiv="re&#102;resh" content="0;url=https://evil.example/?d=S">',
    '<meta content="0;url=https://evil.example/?a>b" http-equiv="refresh">',
    '<img src="x"onerror=alert(1)>',
    '<iframe srcdoc="&lt;script&gt;alert(1)&lt;/script&gt;">x</iframe>',
    '<svg><set attributeName="onload" to="alert(1)"/></svg>',
    '<a href="java&#115;cript:alert(1)">x</a>',
    '<object data="https://evil.example/o">x</object>',
    '<base href="https://evil.example/">',
    '<link rel="stylesheet" href="https://evil.example/c.css">',
    '<img src="https://evil.example/pixel.png">',
    '<form action="https://evil.example/"><button>Go</button></form>',
    '<style>a{}</style ><script>alert(1)</script>',
])
def test_html_sanitizer_blocks_navigation_and_script(payload: str) -> None:
    out = af.sanitize_html(payload).lower()
    for bad in ("evil.example", "refresh", "onerror", "onload", "javascript",
                "srcdoc", "<script", "<iframe", "<object", "<base", "<meta",
                "<link", "<form", "<set", "alert("):
        assert bad not in out, (bad, out)


def test_html_sanitizer_keeps_styling_and_forces_links_to_new_tab() -> None:
    out = af.sanitize_html(
        '<style>h1 > span { color: red }</style><h1 style="color:blue">A &amp; B</h1>'
        '<a href="https://ok.example/x" target="_self">ok</a><a href="#s">jump</a>'
        '<img src="data:image/png;base64,AAAA"><svg><rect width="4" height="2"/></svg>'
    )
    assert "<style>h1 > span { color: red }</style>" in out
    assert '<h1 style="color:blue">A &amp; B</h1>' in out
    assert ('<a href="https://ok.example/x" target="_blank" '
            'rel="noopener noreferrer nofollow">ok</a>') in out
    assert '<a href="#s">jump</a>' in out
    assert 'src="data:image/png;base64,AAAA"' in out
    assert '<rect width="4" height="2" />' in out
