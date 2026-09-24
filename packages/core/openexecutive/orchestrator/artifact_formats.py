"""Artifact formats: how a drafted deliverable is validated, stored and served.

Every artifact is persisted as TEXT in the row that owns it (`alerts.body`
for drafts, `workflow_runs.artifact` for workflow output). Binary formats are
never written to disk — a `.docx` is rendered from the stored Markdown and an
`.xlsx` from the stored sheet JSON at download time. That keeps one source of
truth per artifact, needs no file volume, and means delete/archive have
nothing extra to clean up.

| format   | stored text                         | download         |
|----------|-------------------------------------|------------------|
| markdown | the Markdown                        | .md              |
| html     | the HTML (scripts/refresh stripped) | .html            |
| docx     | the Markdown source                 | .docx (rendered) |
| xlsx     | JSON `{"summary", "sheets"}`        | .xlsx (rendered) |
| link     | Markdown summary (url in own column) | none             |

`link` records a deliverable that lives in one of the principal's connected
apps (any MCP server: a Google Sheet, a Notion page, an Excel Online
workbook…). Nothing here knows about a particular vendor.

`ARTIFACT_FORMAT_NAMES` is a hard-coded, sorted tuple because it feeds the
`draft_artifact` tool schema, which sits in the prompt-cached tool prefix:
its bytes must not depend on dict iteration order.
"""
from __future__ import annotations

import io
import json
import re
from collections.abc import Callable
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Any
from urllib.parse import urlsplit

ARTIFACT_FORMAT_NAMES: tuple[str, ...] = ("docx", "html", "link", "markdown", "xlsx")
DEFAULT_FORMAT = "markdown"

# Largest document (Markdown / HTML) the tool accepts. A drafted artifact is
# echoed into the chat transcript as a tool input, so this also bounds how
# much one draft can grow the conversation.
MAX_DOCUMENT_CHARS = 60_000
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 5_000
MAX_COLUMNS = 200
MAX_TOTAL_CELLS = 200_000
MAX_URL_CHARS = 2_000
MAX_LINK_LABEL_CHARS = 60
# Rows per sheet shown in the on-screen / read-back table preview. The full
# sheet is always in the .xlsx download.
PREVIEW_ROWS = 50
# Excel's own limit on sheet names, and the characters it rejects in them.
_SHEET_NAME_MAX = 31
_SHEET_NAME_BAD = re.compile(r"[\[\]:*?/\\]")


class ArtifactInputError(ValueError):
    """The tool input can't be turned into an artifact of the asked format."""


@dataclass(frozen=True)
class BuiltArtifact:
    """What `build` hands the persistence layer."""

    stored: str
    url: str | None = None
    link_label: str | None = None


@dataclass(frozen=True)
class ArtifactFormat:
    name: str
    label: str
    mime: str
    # File extension for the download, or None when the format has no file.
    extension: str | None
    build: Callable[[dict[str, Any]], BuiltArtifact]
    # Markdown shown on screen / returned by `get_artifact` (html returns its
    # tag-stripped text there; the UI renders the raw HTML in a sandbox).
    display: Callable[[str], str]
    # Plain text for previews and knowledge indexing.
    text: Callable[[str], str]
    # Bytes for the download, or None when the format has no file.
    render_file: Callable[[str], bytes] | None


# --------------------------------------------------------------------------- #
# Shared helpers
# --------------------------------------------------------------------------- #


def _document(tool_input: dict[str, Any], *, required: bool = True) -> str:
    doc = str(tool_input.get("document") or "").strip()
    if required and not doc:
        raise ArtifactInputError("document is required for this format")
    if len(doc) > MAX_DOCUMENT_CHARS:
        raise ArtifactInputError(
            f"document is {len(doc)} chars; the limit is {MAX_DOCUMENT_CHARS}"
        )
    return doc


def with_sources_footer(document: str, source_urls: Any) -> str:
    """Append a Markdown 'Sources' list for any non-blank URLs."""
    if not isinstance(source_urls, list):
        return document
    urls = [str(u).strip() for u in source_urls if str(u).strip()]
    if not urls:
        return document
    footer = "\n\n### Sources\n" + "\n".join(f"- {u}" for u in urls)
    return document + footer


def _collapse(text: str) -> str:
    return " ".join(text.split())


# C0 control characters (bar tab / newline / carriage return) are illegal in
# the XML inside .docx / .xlsx; text scraped from web pages or PDFs often
# carries them. Stripped before rendering so a download never 500s.
_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def xml_safe(text: str) -> str:
    return _XML_ILLEGAL_RE.sub("", text)


def _markdown_text(markdown: str) -> str:
    """Markdown minus a single leading heading line, whitespace-collapsed."""
    text = (markdown or "").lstrip()
    # A lone heading line (no body yet) is kept, so the preview isn't empty.
    if text.startswith("#"):
        first_break = text.find("\n")
        if first_break != -1:
            text = text[first_break + 1 :]
    return _collapse(text)


# --------------------------------------------------------------------------- #
# markdown
# --------------------------------------------------------------------------- #


def _build_markdown(tool_input: dict[str, Any]) -> BuiltArtifact:
    doc = _document(tool_input)
    return BuiltArtifact(stored=with_sources_footer(doc, tool_input.get("source_urls")))


# --------------------------------------------------------------------------- #
# html
# --------------------------------------------------------------------------- #

# Removed from stored HTML. This is best-effort hygiene for the stored copy,
# NOT the security boundary: the viewer renders in a script-less sandboxed
# iframe with a no-network CSP, and downloads are attachments with a sandbox
# CSP header. A downloaded .html opened locally is only as safe as this
# regex pass, so treat it like any file from the web.
_SCRIPT_RE = re.compile(r"<script\b[^>]*>.*?</script\s*>", re.IGNORECASE | re.DOTALL)
_SCRIPT_OPEN_RE = re.compile(r"<script\b[^>]*>", re.IGNORECASE)
_META_REFRESH_RE = re.compile(r"<meta\b[^>]*http-equiv\s*=\s*[\"']?refresh[^>]*>", re.IGNORECASE)
_BASE_RE = re.compile(r"<base\b[^>]*>", re.IGNORECASE)
_TAG_RE = re.compile(r"<[a-zA-Z][^>]*>")
# Only applied INSIDE a tag, so body text like "online = true" is untouched.
_EVENT_ATTR_RE = re.compile(r"[\s/]on[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", re.IGNORECASE)
_JS_URL_RE = re.compile(r"javascript\s*:", re.IGNORECASE)


def _clean_tag(match: re.Match[str]) -> str:
    tag = _EVENT_ATTR_RE.sub(" ", match.group(0))
    return _JS_URL_RE.sub("blocked:", tag)


def sanitize_html(html: str) -> str:
    out = _SCRIPT_RE.sub("", html)
    out = _SCRIPT_OPEN_RE.sub("", out)
    out = _META_REFRESH_RE.sub("", out)
    out = _BASE_RE.sub("", out)
    return _TAG_RE.sub(_clean_tag, out)


class _TextExtractor(HTMLParser):
    _SKIP = frozenset({"script", "style", "head", "title", "noscript", "template"})
    _BLOCK = frozenset({
        "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
        "section", "article", "table", "ul", "ol", "blockquote", "pre",
    })

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    """Readable plain text of an HTML document, paragraphs kept."""
    parser = _TextExtractor()
    try:
        parser.feed(html or "")
        parser.close()
    except Exception:
        return _collapse(re.sub(r"<[^>]+>", " ", html or ""))
    lines = [_collapse(line) for line in "".join(parser.parts).splitlines()]
    return "\n\n".join(line for line in lines if line)


def _build_html(tool_input: dict[str, Any]) -> BuiltArtifact:
    doc = sanitize_html(_document(tool_input))
    if not html_to_text(doc):
        raise ArtifactInputError("html document has no visible text")
    return BuiltArtifact(stored=doc)


# --------------------------------------------------------------------------- #
# docx (stored as Markdown, rendered to Word on download)
# --------------------------------------------------------------------------- #

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_HR_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
# **bold**, __bold__, *italic*, _italic_, `code`, [text](url)
# Underscore emphasis must not touch intraword underscores (snake_case,
# john_doe@example.com), matching CommonMark.
_INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|(?<!\w)__[^_]+__(?!\w)|`[^`]+`|\[[^\]]+\]\([^)]+\)"
    r"|\*[^*\s][^*]*\*|(?<!\w)_[^_\s][^_]*_(?!\w))"
)


def _add_inline(paragraph: Any, text: str) -> None:
    for piece in _INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith(("**", "__")) and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).font.name = "Courier New"
        elif piece.startswith("[") and "](" in piece:
            label, _, url = piece[1:-1].partition("](")
            paragraph.add_run(f"{label} ({url})")
        elif piece[0] in "*_" and piece[-1] == piece[0] and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _table_cells(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [c.strip() for c in row.split("|")]


def _is_table_start(lines: list[str], i: int) -> bool:
    """GFM: a header row containing '|' and a delimiter row of equal width."""
    if "|" not in lines[i] or i + 1 >= len(lines):
        return False
    sep = lines[i + 1]
    if "|" not in sep or not _TABLE_SEP_RE.match(sep):
        return False
    return len(_table_cells(sep)) == len(_table_cells(lines[i]))


def _emit_code_block(doc: Any, lines: list[str], i: int) -> int:
    i += 1
    code: list[str] = []
    while i < len(lines) and not lines[i].strip().startswith("```"):
        code.append(lines[i])
        i += 1
    doc.add_paragraph().add_run("\n".join(code)).font.name = "Courier New"
    return i + 1


def _emit_table(doc: Any, lines: list[str], i: int) -> int:
    header = _table_cells(lines[i])
    i += 2
    body: list[list[str]] = []
    while i < len(lines) and "|" in lines[i] and lines[i].strip():
        body.append(_table_cells(lines[i]))
        i += 1
    # A body row wider than the header widens the table rather than losing cells.
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    table = doc.add_table(rows=1 + len(body), cols=width)
    table.style = "Table Grid"
    for c in range(width):
        cell_text = header[c] if c < len(header) else ""
        table.rows[0].cells[c].paragraphs[0].add_run(cell_text).bold = True
    for r, row in enumerate(body, start=1):
        for c in range(width):
            _add_inline(table.rows[r].cells[c].paragraphs[0], row[c] if c < len(row) else "")
    return i


def _emit_list_item(doc: Any, line: str) -> bool:
    bullet = _BULLET_RE.match(line)
    match = bullet or _NUMBERED_RE.match(line)
    if match is None:
        return False
    nested = len(match.group(1).expandtabs(4)) >= 2
    style = "List Bullet" if bullet else "List Number"
    _add_inline(doc.add_paragraph(style=f"{style} 2" if nested else style), match.group(2).strip())
    return True


def markdown_to_docx(markdown: str) -> bytes:
    """Render the Markdown subset artifacts use into a Word document.

    Covers headings, paragraphs, bullet / numbered lists (one nesting level),
    GFM tables, block quotes, fenced code, rules and inline bold / italic /
    code / links. Anything else is kept as literal paragraph text, so no
    content is dropped (control characters Word can't store aside).
    """
    from docx import Document

    doc = Document()
    lines = xml_safe(markdown or "").splitlines()
    i = 0
    paragraph_buf: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buf:
            _add_inline(doc.add_paragraph(), " ".join(s.strip() for s in paragraph_buf))
            paragraph_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        heading = _HEADING_RE.match(stripped)
        if stripped and not (stripped.startswith(("```", ">")) or heading
                             or _is_table_start(lines, i) or _HR_RE.match(stripped)
                             or _BULLET_RE.match(line) or _NUMBERED_RE.match(line)):
            paragraph_buf.append(line)
            i += 1
            continue
        flush_paragraph()
        if not stripped:
            i += 1
        elif stripped.startswith("```"):
            i = _emit_code_block(doc, lines, i)
        elif heading:
            _add_inline(doc.add_heading(level=min(len(heading.group(1)), 4)),
                        heading.group(2).strip())
            i += 1
        elif _is_table_start(lines, i):
            i = _emit_table(doc, lines, i)
        elif _HR_RE.match(stripped):
            doc.add_paragraph("―" * 20)
            i += 1
        elif _emit_list_item(doc, line):
            i += 1
        else:  # block quote
            _add_inline(doc.add_paragraph(style="Intense Quote"), stripped.lstrip("> ").strip())
            i += 1

    flush_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# xlsx (stored as sheet JSON, rendered to Excel on download)
# --------------------------------------------------------------------------- #

Cell = str | int | float | bool | None


def _cell(value: Any) -> Cell:
    if isinstance(value, str):
        return xml_safe(value)
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return xml_safe(json.dumps(value, ensure_ascii=False))


def _sheet_name(raw: Any, index: int, seen: set[str]) -> str:
    name = _SHEET_NAME_BAD.sub(" ", xml_safe(str(raw or "")).strip())[:_SHEET_NAME_MAX].strip()
    name = name or f"Sheet{index + 1}"
    base, n = name, 2
    while name.lower() in seen:
        suffix = f" ({n})"
        name = base[: _SHEET_NAME_MAX - len(suffix)] + suffix
        n += 1
    seen.add(name.lower())
    return name


def _normalize_sheets(raw: Any) -> list[dict[str, Any]]:
    if not isinstance(raw, list) or not raw:
        raise ArtifactInputError(
            "sheets is required for xlsx: a list of {name, columns, rows}"
        )
    if len(raw) > MAX_SHEETS:
        raise ArtifactInputError(f"at most {MAX_SHEETS} sheets")
    seen: set[str] = set()
    sheets: list[dict[str, Any]] = []
    total_cells = 0
    for idx, sheet in enumerate(raw):
        if not isinstance(sheet, dict):
            raise ArtifactInputError(f"sheet {idx + 1} must be an object")
        columns = sheet.get("columns") or []
        rows = sheet.get("rows") or []
        if not isinstance(columns, list) or not isinstance(rows, list):
            raise ArtifactInputError(f"sheet {idx + 1}: columns and rows must be lists")
        if not columns and not rows:
            raise ArtifactInputError(f"sheet {idx + 1} is empty")
        if len(columns) > MAX_COLUMNS:
            raise ArtifactInputError(f"sheet {idx + 1}: at most {MAX_COLUMNS} columns")
        if len(rows) > MAX_ROWS_PER_SHEET:
            raise ArtifactInputError(
                f"sheet {idx + 1}: at most {MAX_ROWS_PER_SHEET} rows"
            )
        # Models often send rows as {column: value} objects; map them onto
        # the columns (declaring columns from the first row if none given).
        if not columns and rows and isinstance(rows[0], dict):
            columns = list(rows[0].keys())
        norm_rows: list[list[Cell]] = []
        for row in rows:
            if isinstance(row, dict):
                cells = [row.get(c) for c in columns]
            else:
                cells = row if isinstance(row, list) else [row]
            if len(cells) > MAX_COLUMNS:
                raise ArtifactInputError(f"sheet {idx + 1}: at most {MAX_COLUMNS} columns")
            norm_rows.append([_cell(v) for v in cells])
            total_cells += len(cells)
        if total_cells > MAX_TOTAL_CELLS:
            raise ArtifactInputError(f"at most {MAX_TOTAL_CELLS} cells in total")
        sheets.append({
            "name": _sheet_name(sheet.get("name"), idx, seen),
            "columns": [xml_safe(str(c)) for c in columns],
            "rows": norm_rows,
        })
    return sheets


def _build_xlsx(tool_input: dict[str, Any]) -> BuiltArtifact:
    summary = _document(tool_input, required=False)
    summary = with_sources_footer(summary, tool_input.get("source_urls")).strip()
    sheets = _normalize_sheets(tool_input.get("sheets"))
    stored = json.dumps({"summary": summary, "sheets": sheets}, ensure_ascii=False)
    return BuiltArtifact(stored=stored)


def _load_workbook_json(stored: str) -> dict[str, Any]:
    try:
        data = json.loads(stored or "{}")
    except json.JSONDecodeError:
        return {"summary": stored or "", "sheets": []}
    if not isinstance(data, dict):
        return {"summary": "", "sheets": []}
    return data


def _md_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ")


def _xlsx_display(stored: str) -> str:
    data = _load_workbook_json(stored)
    parts: list[str] = []
    if data.get("summary"):
        parts.append(str(data["summary"]))
    for sheet in data.get("sheets") or []:
        columns = sheet.get("columns") or []
        rows = sheet.get("rows") or []
        width = max(1, len(columns), *(len(r) for r in rows))
        header = [*columns, *[""] * (width - len(columns))]
        lines = [f"### {sheet.get('name', 'Sheet')}", ""]
        lines.append("| " + " | ".join(_md_cell(c) for c in header) + " |")
        lines.append("|" + "---|" * width)
        for row in rows[:PREVIEW_ROWS]:
            padded = [*row, *[""] * (width - len(row))]
            lines.append("| " + " | ".join(_md_cell(c) for c in padded) + " |")
        if len(rows) > PREVIEW_ROWS:
            lines.append("")
            lines.append(
                f"_Showing {PREVIEW_ROWS} of {len(rows)} rows — download the "
                "workbook for the rest._"
            )
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _xlsx_text(stored: str) -> str:
    data = _load_workbook_json(stored)
    bits = [str(data.get("summary") or "")]
    for sheet in data.get("sheets") or []:
        cols = ", ".join(str(c) for c in sheet.get("columns") or [])
        bits.append(f"{sheet.get('name', 'Sheet')}: {len(sheet.get('rows') or [])} rows ({cols})")
    return _collapse(" ".join(b for b in bits if b))


def sheets_to_xlsx(stored: str) -> bytes:
    """Render stored sheet JSON to a workbook.

    Every string is written as a literal string, never as a formula: sheet
    content can carry text quoted from the web or email, and a cell that
    starts with '=' must not execute in the reader's spreadsheet app.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    data = _load_workbook_json(stored)
    wb = Workbook()
    default = wb.active
    sheets = data.get("sheets") or []
    if default is not None and sheets:
        wb.remove(default)
    seen: set[str] = set()
    for sheet in sheets:
        ws = wb.create_sheet(title=_sheet_name(sheet.get("name"), len(seen), seen))
        columns = sheet.get("columns") or []
        rows = sheet.get("rows") or []
        widths: dict[int, int] = {}
        start = 1
        if columns:
            for c, name in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=c, value=xml_safe(str(name)))
                cell.data_type = "s"
                cell.font = Font(bold=True)
                widths[c] = max(widths.get(c, 0), len(str(name)))
            ws.freeze_panes = "A2"
            start = 2
        for r, row in enumerate(rows, start=start):
            for c, value in enumerate(row, start=1):
                value = _cell(value)
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                widths[c] = max(widths.get(c, 0), len("" if value is None else str(value)))
        for c, width in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# link (a deliverable living in a connected app)
# --------------------------------------------------------------------------- #


def validate_link_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        raise ArtifactInputError("url is required for a link artifact")
    if len(url) > MAX_URL_CHARS:
        raise ArtifactInputError("url is too long")
    try:
        parts = urlsplit(url)
        host = parts.hostname
    except ValueError as exc:
        raise ArtifactInputError(f"url is not valid: {exc}") from exc
    if parts.scheme != "https" or not host:
        raise ArtifactInputError("url must be an https:// link")
    if parts.username or parts.password:
        raise ArtifactInputError("url must not embed credentials")
    if any(ch.isspace() for ch in url):
        raise ArtifactInputError("url must not contain whitespace")
    return url


def _build_link(tool_input: dict[str, Any]) -> BuiltArtifact:
    url = validate_link_url(str(tool_input.get("url") or ""))
    label = _collapse(str(tool_input.get("link_label") or ""))[:MAX_LINK_LABEL_CHARS]
    summary = _document(tool_input, required=False)
    summary = with_sources_footer(summary, tool_input.get("source_urls"))
    return BuiltArtifact(stored=summary, url=url, link_label=label or "External document")


# --------------------------------------------------------------------------- #
# Registry
# --------------------------------------------------------------------------- #


def _utf8(text: str) -> bytes:
    return (text or "").encode("utf-8")


ARTIFACT_FORMATS: dict[str, ArtifactFormat] = {
    "docx": ArtifactFormat(
        name="docx", label="Word document",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension="docx", build=_build_markdown, display=lambda s: s,
        text=_markdown_text, render_file=markdown_to_docx,
    ),
    "html": ArtifactFormat(
        name="html", label="Web page", mime="text/html; charset=utf-8",
        extension="html", build=_build_html, display=html_to_text,
        text=lambda s: _collapse(html_to_text(s)), render_file=_utf8,
    ),
    "link": ArtifactFormat(
        name="link", label="Link", mime="text/plain; charset=utf-8",
        extension=None, build=_build_link, display=lambda s: s,
        text=_markdown_text, render_file=None,
    ),
    "markdown": ArtifactFormat(
        name="markdown", label="Document", mime="text/markdown; charset=utf-8",
        extension="md", build=_build_markdown, display=lambda s: s,
        text=_markdown_text, render_file=_utf8,
    ),
    "xlsx": ArtifactFormat(
        name="xlsx", label="Spreadsheet",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        extension="xlsx", build=_build_xlsx, display=_xlsx_display,
        text=_xlsx_text, render_file=sheets_to_xlsx,
    ),
}

# Extra download targets a stored format can be converted to on request
# (e.g. any Markdown artifact — including every workflow output — as Word).
EXPORT_TARGETS: dict[str, tuple[str, ...]] = {
    "markdown": ("docx",),
    "docx": ("markdown",),
}


def get_format(name: str | None) -> ArtifactFormat:
    """The registered format, falling back to Markdown for unknown / legacy rows."""
    return ARTIFACT_FORMATS.get((name or DEFAULT_FORMAT).strip().lower(),
                                ARTIFACT_FORMATS[DEFAULT_FORMAT])


def build_artifact(fmt_name: str, tool_input: dict[str, Any]) -> BuiltArtifact:
    name = (fmt_name or DEFAULT_FORMAT).strip().lower()
    if name not in ARTIFACT_FORMATS:
        raise ArtifactInputError(
            f"unknown format {fmt_name!r}; use one of {', '.join(ARTIFACT_FORMAT_NAMES)}"
        )
    return ARTIFACT_FORMATS[name].build(tool_input)


def preview_text(fmt_name: str | None, stored: str, n: int) -> str:
    return get_format(fmt_name).text(stored)[:n]


__all__ = [
    "ARTIFACT_FORMATS",
    "ARTIFACT_FORMAT_NAMES",
    "DEFAULT_FORMAT",
    "EXPORT_TARGETS",
    "ArtifactFormat",
    "ArtifactInputError",
    "BuiltArtifact",
    "build_artifact",
    "get_format",
    "html_to_text",
    "markdown_to_docx",
    "preview_text",
    "sanitize_html",
    "sheets_to_xlsx",
    "validate_link_url",
    "with_sources_footer",
    "xml_safe",
]
